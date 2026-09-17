"""
Maps C10 reader and imager files (.xlsx or .txt) to measurements.

The file name has the date, the time and the plate barcode, e.g.
"241014_125455_241008MP-1_1.txt". A file starts with metadata, followed by a
"Results" table with one row per well. An .xlsx file can also have a "Layout"
block that marks positive (POS) and negative (NEG) control wells.
"""

import os
import re
from typing import TypedDict

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from tqdm import tqdm

from core.models import Measurement, Well, WellType
from helpers.logger import logger
from importer.helper import message
from importer.mappers.base import BaseMapper
from importer.mappers.values import convert_sci_to_float, parse_c10_datetime

# The label of the values of a .txt file that has no header line above its values
DEFAULT_MEASUREMENT_NAME = "Lum"

# Only this many rows at the top of an .xlsx sheet are read as metadata
XLSX_METADATA_ROWS = 38

# Words in C10 files that start a block
RESULTS_BLOCK = "Results"
LAYOUT_BLOCK = "Layout"

# The keys of the measurement date and time in the metadata of a .txt file
TXT_DATE_KEY = "Date"
TXT_TIME_KEY = "Time"

# Result columns that name the well; they are not measurements
WELL_ID_COLUMN = "Well ID"
WELL_COLUMN = "Well"
WELL_COLUMNS = (WELL_ID_COLUMN, WELL_COLUMN)

# Values in the "Layout" block -> the well type names they stand for
LAYOUT_WELL_TYPES = {"POS": "P", "NEG": "N"}

# A measurement value is a plain number, e.g. "16727" or "1.5" ...
RE_NUMBER = r"^[0-9]+(\.[0-9]+)?$"
# ... or a number in scientific notation, e.g. "1.5E+03"
RE_SCIENCE = r"^[0-9\.]+[eE][+-]?[0-9]+$"


class C10Data(TypedDict):
    """What `parse` returns for one C10 file."""

    barcode: str  # e.g. "241008MP-1_1"
    date: str  # from the .txt content ("10/14/2024") or the file name ("241014")
    time: str  # e.g. "12:45:28" or "125455"
    metadata: dict  # the lines above the results, e.g. {"Plate Number": "Plate 1"}
    results: list[dict]  # one dict per well, e.g. {"Well": "A1", "Lum": "16727"}
    layout: dict[str, str]  # control wells, e.g. {"A1": "P"}; empty for .txt files


class MicroscopeMapper(BaseMapper):
    # e.g. "241014_125455_241008MP-1_1.txt": date, time, barcode and extension
    RE_FILENAME = (
        r"(?P<date>\d+)[-_](?P<time>\d+)[-_](?P<barcode>[^\.]+)\.(?P<ext>xlsx|txt)$"
    )

    def parse(self, file: str, **kwargs) -> C10Data:
        """
        Reads a C10 file; `file` is the file name.

        Returned data example:
        {"barcode": "241008MP-1_1", "date": "10/14/2024", "time": "12:45:28",
         "metadata": {"Plate Number": "Plate 1", ...},
         "results": [{"Well": "A1", "Lum": "16727"}, ...],
         "layout": {"A1": "P", "B1": "N"}}
        """
        date, time, barcode, extension = self.file_name_parts(
            file, kwargs.get("room_name")
        )
        # None or "" when no name is given on the management page
        measurement_name = kwargs.get("measurement_name")

        if extension == "xlsx":
            sheet = load_workbook(file).active
            metadata = self.parse_xlsx_metadata(sheet)
            results = self.parse_xlsx_results(sheet)
            layout = self.parse_xlsx_layout(sheet)
        elif extension == "txt":
            with open(file, "r") as content:
                lines = [line.strip() for line in content.readlines() if line.strip()]
            metadata = self.parse_txt_metadata(lines)
            results = self.parse_txt_results(lines, measurement_name)
            # A .txt file has no layout block
            layout = {}
            # The date and time in the file content win over the file name
            date = metadata.get(TXT_DATE_KEY, date)
            time = metadata.get(TXT_TIME_KEY, time)
            logger.info(f"Date: {date}, Time: {time}")
        else:
            raise ValueError(f"Unsupported file extension: {extension}")

        return {
            "metadata": metadata,
            "results": results,
            "date": date,
            "time": time,
            "barcode": barcode,
            "layout": layout,
        }

    def file_name_parts(
        self, path: str, room_name: str | None
    ) -> tuple[str, str, str, str]:
        """
        "/data/241014_125455_241008MP-1_1.txt" -> ("241014", "125455", "241008MP-1_1", "txt").
        """
        file_name = os.path.basename(path)
        match = re.match(self.RE_FILENAME, file_name)
        if not match:
            text = f"Filename {file_name} does not match expected pattern."
            message(text, "error", room_name)
            raise ValueError(text)
        return (
            match.group("date"),
            match.group("time"),
            match.group("barcode"),
            match.group("ext"),
        )

    def map(self, data: C10Data, **kwargs) -> None:
        """
        Stores every number of the results as a measurement of its well, sets
        the control well types from the layout, and links the file to the plate.

        A file with an unknown date format stops the mapping before anything
        of it is stored.
        """
        measured_at = parse_c10_datetime(data["date"], data["time"])
        if measured_at is None:
            # The map command shows this error on the management page
            raise ValueError(
                f"Cannot read the measurement date of {kwargs.get('filename')}: "
                f"date '{data['date']}', time '{data['time']}'. "
                "Nothing of this file was stored, and the next files were not mapped."
            )

        plate = self.find_or_create_measured_plate(
            data["barcode"],
            len(data["results"]),
            kwargs.get("room_name"),
            kwargs.get("experiment_name"),
        )

        with tqdm(
            desc="Processing microscope output",
            unit="measurement",
            total=len(data["results"]),
        ) as progress:
            for entry in data["results"]:
                well_name = entry.get(WELL_COLUMN)
                # Empty rows and repeated header rows are not wells
                if not well_name or well_name == WELL_COLUMN:
                    continue

                position = plate.dimension.position(well_name)
                well = plate.well_at(position, create_if_not_exist=True)
                self.set_well_type_from_layout(well, well_name, data["layout"])

                for label, value in entry.items():
                    if label in WELL_COLUMNS:
                        continue
                    if re.match(RE_NUMBER, str(value)):
                        value = float(value)
                    elif re.match(RE_SCIENCE, str(value)):
                        value = convert_sci_to_float(value)
                    else:
                        # Not a number, e.g. a text column
                        continue
                    Measurement.objects.update_or_create(
                        well=well,
                        label=label,
                        measured_at=measured_at,
                        defaults={"value": value},
                    )
                progress.update(1)

        self.create_measurement_assignment(plate, kwargs.get("filename"))

    @staticmethod
    def set_well_type_from_layout(
        well: Well, well_name: str, layout: dict[str, str]
    ) -> None:
        """Sets the well type if the layout has the well, e.g. {"A1": "P"}."""
        if layout and well_name in layout:
            well.type = WellType.objects.get(name=layout[well_name])
            well.save()

    @staticmethod
    def parse_xlsx_metadata(sheet: Worksheet) -> dict:
        """
        The metadata at the top of the sheet. A text in the first column starts
        a label, and the other cells of the row are added to that label. A cell
        like "Gain: 214" is stored under its own key instead.

        Example: {"Plate Number": ["Plate 1"], "Read": ["Luminescence"], "Gain": "214"}
        """
        # Values are text or lists of text, see the example above
        metadata: dict = {}
        current_label = None
        for row in sheet.iter_rows(min_row=1, max_row=XLSX_METADATA_ROWS):
            for index, cell in enumerate(row):
                if index == 0 and cell.value:
                    current_label = str(cell.value).strip().replace(":", "")
                    metadata[current_label] = []
                elif index != 0 and cell.value:
                    text = str(cell.value)
                    if ": " in text:
                        # The value is the part between the first and a second ":"
                        key = text.split(":")[0].strip()
                        value = text.split(":")[1].strip()
                        metadata[key] = value
                    else:
                        metadata[current_label].append(text)
                # "Results" ends this row; the next rows are still read
                if str(cell.value).strip().lower() == RESULTS_BLOCK.lower():
                    break
        return metadata

    @staticmethod
    def parse_xlsx_results(sheet: Worksheet) -> list[dict]:
        """
        The rows below the header row of the results table, which has "Well ID"
        or "Well" in its second column.

        Example: [{"Well ID": "SPL1", "Well": "A1", "Lum": 16727}]
        """
        header_rows = []
        results_start_row = None
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            if str(row[1]).lower() in (WELL_ID_COLUMN.lower(), WELL_COLUMN.lower()):
                header_rows.append(row)
                # Excel rows count from 1, and the results start below the header
                results_start_row = index + 2
                break

        # A sheet without header row stops here with an IndexError
        headers = [header for header in header_rows[0] if header is not None]
        results = []
        if headers and results_start_row:
            for row in sheet.iter_rows(min_row=results_start_row, values_only=True):
                if not any(row):
                    continue
                # The first column of a result row is empty
                results.append(dict(zip(headers, row[1:])))
        return results

    @staticmethod
    def parse_xlsx_layout(sheet: Worksheet) -> dict[str, str]:
        """
        The control wells of the "Layout" block, which starts one row below
        "Layout" and ends two rows above "Results". A layout row starts with
        the row letter, followed by one cell per column.

        Example: {"A1": "P", "B1": "N"}
        """
        layout_start_row = None
        layout_end_row = None
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            if row[0] and str(row[0]).lower() == LAYOUT_BLOCK.lower():
                layout_start_row = index + 2
            if row[0] and str(row[0]).lower() == RESULTS_BLOCK.lower():
                layout_end_row = index - 1

        layout_rows = []
        if layout_start_row and layout_end_row:
            for row in sheet.iter_rows(
                min_row=layout_start_row, max_row=layout_end_row, values_only=True
            ):
                if not any(row):
                    continue
                layout_rows.append(row[1:])

        well_types = {}
        for layout_row in layout_rows:
            row_letter = layout_row[0]
            if not row_letter:
                continue
            for column, value in enumerate(layout_row):
                if value in LAYOUT_WELL_TYPES:
                    well_types[f"{row_letter}{column}"] = LAYOUT_WELL_TYPES[value]
        return well_types

    @staticmethod
    def parse_txt_metadata(lines: list[str]) -> dict:
        """
        The "Key<TAB>Value" and "Key: Value" lines before "Results".

        Example: {"Date": "10/14/2024", "Integration Time": "0:01.00 (MM:SS.ss)"}
        """
        metadata = {}
        for line in lines:
            line = line.strip()
            if line == RESULTS_BLOCK:
                break
            if not line:
                continue
            if "\t" in line:
                key, value = line.split("\t", 1)
                metadata[key.strip()] = value.strip()
            elif ":" in line:
                key, value = line.split(":", 1)
                metadata[key.strip()] = value.strip()
        return metadata

    @staticmethod
    def parse_txt_results(lines: list[str], measurement_name: str | None) -> list[dict]:
        """
        The "Well<TAB>Value" lines after "Results". The values are labeled with
        the given measurement name, or else with the name in the header line
        ("Well<TAB>Lum").

        Example: [{"Well": "A1", "Lum": "16727"}]
        """
        label = measurement_name or DEFAULT_MEASUREMENT_NAME
        results = []
        in_results = False
        for line in lines:
            if not in_results:
                if line == RESULTS_BLOCK:
                    in_results = True
                continue

            if line.strip() == "":
                break
            parts = line.split("\t")
            if len(parts) != 2:
                continue
            well_name, value = parts
            # The header line of the table
            if well_name.strip() == WELL_COLUMN:
                if not measurement_name and value.strip():
                    label = value.strip()
                continue
            results.append({WELL_COLUMN: well_name.strip(), label: value.strip()})
        return results
